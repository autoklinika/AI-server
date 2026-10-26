"""Bounded, explicit spreadsheet formats. No user/session tables are exported."""
from __future__ import annotations

import base64
import binascii
import csv
import hashlib
import hmac
import io
import json
import re
import time
import unicodedata
import zipfile
from xml.etree import ElementTree

from .store import KINDS, PRODUCT_STATES, TECH_STATES, Problem, normalize

MAX_FILE = 1_000_000
MAX_ROWS = 1000
MAX_BODY = 1_400_000
FIELDS = ("kind", "brand", "designation", "manufacturer", "category", "model",
          "primary_number", "hw", "sw", "package", "parameters", "notes")
IMPORT_HEADERS = ("format", *FIELDS, "brands_json", "numbers_json", "location",
                  "product_state", "technical_state", "quantity")
EXPORT_HEADERS = ("format", "item_id", *FIELDS, "brands_json", "numbers_json", "location",
                  "product_state", "technical_state", "quantity")


def spreadsheet_library():
    try:
        import openpyxl
    except ImportError:
        raise Problem(503, "XLSX wymaga opcjonalnej zależności apps/magazyn/requirements-xlsx.txt")
    return openpyxl


def safe_cell(value):
    if not isinstance(value, str):
        return value
    # Excel can discard whitespace/control prefixes before interpreting formulas.
    probe = value.lstrip()
    probe = ''.join(c for c in probe if not unicodedata.category(c).startswith('C')).lstrip()
    if probe.startswith(('=', '+', '-', '@')):
        value = "'" + value
    # XML cannot represent most control characters; preserve them visibly.
    return ''.join(c if c in '\t\n\r' or ord(c) >= 32 else '\\x%02x' % ord(c) for c in value)


def encode_table(headers, rows, fmt):
    if fmt == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows([safe_cell(v) for v in row] for row in rows)
        return stream.getvalue().encode("utf-8-sig")
    if fmt != "xlsx":
        raise Problem(422, "Format musi być csv lub xlsx")
    library = spreadsheet_library()
    book = library.Workbook(write_only=True)
    sheet = book.create_sheet("Magazyn")
    for row in [headers, *rows]:
        cells = []
        for value in row:
            cell = library.cell.WriteOnlyCell(sheet, value=safe_cell(value))
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def export_catalog(store, fmt, kind=None, brand=None):
    if kind is not None and kind not in KINDS:
        raise Problem(422, "Nieprawidłowy rodzaj")
    rows = []
    with store.connect() as con:
        con.execute("BEGIN")
        for item in con.execute("SELECT * FROM items ORDER BY id"):
            brands = [r[0] for r in con.execute("SELECT brand FROM item_brands WHERE item_id=? ORDER BY brand", (item['id'],))]
            if (kind and item['kind'] != kind) or (brand and brand not in brands):
                continue
            numbers = [dict(r) for r in con.execute("SELECT number,number_type FROM item_numbers WHERE item_id=? ORDER BY normalized", (item['id'],))]
            prefix = ["magazyn-snapshot-v1", item['id'], *[item[f] for f in FIELDS],
                      json.dumps(brands, ensure_ascii=False), json.dumps(numbers, ensure_ascii=False)]
            balances = con.execute("SELECT l.code,b.product_state,b.technical_state,b.quantity FROM balances b JOIN locations l ON l.id=b.location_id WHERE item_id=? ORDER BY l.code,b.product_state,b.technical_state", (item['id'],)).fetchall()
            for balance in balances or [("", "", "", 0)]:
                rows.append([*prefix, *balance])
    return encode_table(EXPORT_HEADERS, rows, fmt)


def decode_upload(content):
    if len(content) > (MAX_FILE + 2) // 3 * 4:
        raise Problem(413, "Plik przekracza 1 MB")
    try:
        raw = base64.b64decode(content, validate=True)
    except (ValueError, binascii.Error):
        raise Problem(422, "Nieprawidłowe base64")
    if len(raw) > MAX_FILE:
        raise Problem(413, "Plik przekracza 1 MB")
    return raw


def xlsx_rows(raw):
    library = spreadsheet_library()
    # Bound decompression before openpyxl; reject active/external content in all parts.
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(e.file_size for e in entries) > 8_000_000:
                raise Problem(413, "Zbyt duży rozpakowany XLSX")
            if len({e.filename for e in entries}) != len(entries):
                raise Problem(422, "Powtórzone części XLSX")
            for entry in entries:
                name = entry.filename.lower()
                if any(x in name for x in ('vba', 'externallink', 'embedding', 'activex')):
                    raise Problem(422, "Makra i linki XLSX są zabronione")
                if name.endswith(('.xml', '.rels')):
                    data = archive.read(entry)
                    if b'<!doctype' in data.lower() or b'<!entity' in data.lower() or b'\x00' in data:
                        raise Problem(422, "Niedozwolony XML")
                    root = ElementTree.fromstring(data)
                    cell_count = row_count = 0
                    for node in root.iter():
                        tag = node.tag.rsplit('}', 1)[-1]
                        if tag in ('f', 'hyperlink', 'oleObject') or node.get('TargetMode', '').lower() == 'external':
                            raise Problem(422, "Formuły i linki XLSX są zabronione")
                        if tag == 'Override' and 'macroenabled' in node.get('ContentType', '').lower():
                            raise Problem(422, "Makra XLSX są zabronione")
                        if tag == 'row' and not 1 <= int(node.get('r', '1')) <= MAX_ROWS + 1:
                            raise Problem(413, "Indeks wiersza poza limitem")
                        if tag == 'c':
                            coord = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', node.get('r', ''))
                            col = 0
                            if coord:
                                for char in coord[1]:
                                    col = col * 26 + ord(char) - 64
                            if not coord or col > len(IMPORT_HEADERS):
                                raise Problem(422, "Niedozwolona kolumna w pliku XLSX")
                            if int(coord[2]) > MAX_ROWS + 1:
                                raise Problem(413, "Zbyt wiele wierszy XLSX")
                        cell_count += tag == 'c'
                        row_count += tag == 'row'
                    if row_count > MAX_ROWS + 1 or cell_count > (MAX_ROWS + 1) * len(IMPORT_HEADERS):
                        raise Problem(413, "Zbyt wiele wierszy lub komórek XLSX")
        book = library.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise Problem(422, "Import wymaga jednego arkusza")
            sheet = book.worksheets[0]
            sheet.reset_dimensions()  # Ignore forged dimensions; XML counts are bounded above.
            for row in sheet.iter_rows():
                if len(row) > len(IMPORT_HEADERS):
                    raise Problem(422, "Za dużo kolumn XLSX")
                yield ["" if c.value is None else str(c.value) for c in row]
        finally:
            book.close()
    except Problem:
        raise
    except Exception as exc:
        raise Problem(422, "Nieprawidłowy XLSX") from exc


def parse_upload(raw, fmt):
    try:
        if fmt == 'csv':
            source = csv.reader(io.StringIO(raw.decode('utf-8-sig'), newline=''), strict=True)
        elif fmt == 'xlsx':
            source = xlsx_rows(raw)
        else:
            raise Problem(422, "Format musi być csv lub xlsx")
        if next(source, None) != list(IMPORT_HEADERS):
            raise Problem(422, "Nieprawidłowe nagłówki: użyj szablonu importu, nie eksportu snapshotu")
        rows = []
        for row in source:
            if len(rows) >= MAX_ROWS:
                raise Problem(413, "Limit importu: 1000 wierszy")
            if len(row) != len(IMPORT_HEADERS):
                raise Problem(422, "Nieprawidłowa liczba kolumn")
            if any(len(value) > 4000 for value in row):
                raise Problem(422, "Zbyt długa komórka")
            rows.append(dict(zip(IMPORT_HEADERS, row)))
        if not rows:
            raise Problem(422, "Pusty plik importu")
        return rows
    except (UnicodeError, csv.Error) as exc:
        raise Problem(422, "Nieprawidłowy CSV UTF-8") from exc


def identity(item):
    return tuple(normalize(item[f]) for f in ('kind', 'designation', 'manufacturer', 'model'))


def validate_rows(con, rows):
    problems, prepared = [], []
    identifiers = {r[0]: 'baza' for r in con.execute('SELECT normalized FROM item_numbers')}
    identities = {identity(dict(r)): 'baza' for r in con.execute('SELECT * FROM items')}
    locations = {r['code'].casefold(): r['id'] for r in con.execute('SELECT * FROM locations')}
    for index, row in enumerate(rows, 2):
        try:
            data = {field: row[field].strip() for field in FIELDS}
            if row['format'] != 'magazyn-new-v1':
                raise ValueError('Wymagany format magazyn-new-v1')
            if data['kind'] not in KINDS or not data['designation'] or (data['kind'] == 'ecu' and not data['brand']):
                raise ValueError('Wymagane kind, designation i marka ECU')
            if any(len(v) > 2000 for v in data.values()):
                raise ValueError('Pole produktu przekracza 2000 znaków')
            brands = json.loads(row['brands_json'] or '[]')
            numbers = json.loads(row['numbers_json'] or '[]')
            if not isinstance(brands, list) or len(brands) > 30 or any(not isinstance(b, str) or not b.strip() or len(b) > 160 for b in brands):
                raise ValueError('Nieprawidłowe brands_json')
            if not isinstance(numbers, list) or len(numbers) > 30:
                raise ValueError('Nieprawidłowe numbers_json')
            pairs = [(data['primary_number'], 'main')] if data['primary_number'] else []
            for n in numbers:
                if not isinstance(n, dict) or set(n) != {'number', 'number_type'} or not all(isinstance(v, str) for v in n.values()):
                    raise ValueError('numbers_json wymaga number i number_type')
                if not n['number_type'] or len(n['number_type']) > 30:
                    raise ValueError('Nieprawidłowy typ numeru')
                pairs.append((n['number'], n['number_type']))
            keys = []
            for number, _ in pairs:
                if not normalize(number) or len(number) > 160:
                    raise ValueError('Nieprawidłowy numer (1–160 znaków)')
                keys.append(normalize(number))
            if len(set(keys)) != len(keys):
                raise ValueError('Powtórzony numer główny/OEM/alternatywny w pozycji')
            # Conservatively block shared identifiers even with different HW/SW.
            conflicts = [identifiers[k] for k in keys if k in identifiers]
            key = identity(data)
            if key in identities:
                conflicts.append(identities[key])
            if conflicts:
                raise ValueError('Konflikt produktu/numeru lub wariantu HW/SW: ' + ', '.join(sorted(set(conflicts))))
            if not re.fullmatch(r'0|[1-9][0-9]{0,9}', row['quantity']) or int(row['quantity']) > 2_000_000_000:
                raise ValueError('Ilość musi być całkowita 0–2000000000')
            location_id = locations.get(row['location'].strip().casefold())
            if location_id is None:
                raise ValueError('Nieznana lokalizacja; utwórz ją przed importem')
            ps, ts = row['product_state'], row['technical_state']
            if data['kind'] == 'ecu':
                if ps not in PRODUCT_STATES or ts not in TECH_STATES:
                    raise ValueError('Nieprawidłowy stan ECU')
            elif (ps, ts) != ('na', 'na'):
                raise ValueError('Elektronika wymaga stanów na/na')
            data.update(brands=brands, numbers=numbers)
            prepared.append((data, dict(location_id=location_id, product_state=ps,
                technical_state=ts, quantity=int(row['quantity']))))
            for number in keys:
                identifiers[number] = f'wiersz {index}'
            identities[key] = f'wiersz {index}'
        except (ValueError, TypeError) as exc:
            problems.append(dict(row=index, detail=str(exc)))
    return prepared, problems


def preview_token(secret, raw, fmt, actor_id, expiry):
    message = f'{actor_id}:{fmt}:{expiry}:'.encode() + hashlib.sha256(raw).digest()
    return f'{expiry}.' + hmac.new(secret, message, hashlib.sha256).hexdigest()


def import_preview(store, raw, fmt, actor, secret):
    rows = parse_upload(raw, fmt)
    with store.connect() as con:
        con.execute('BEGIN')
        prepared, problems = validate_rows(con, rows)
    return dict(count=len(rows), valid_count=len(prepared), problems=problems,
                token=None if problems else preview_token(secret, raw, fmt, actor['id'], int(time.time()) + 900))


def import_apply(store, raw, fmt, actor, secret, token):
    try:
        expiry = int(token.split('.')[0])
    except (ValueError, AttributeError):
        raise Problem(409, 'Najpierw wykonaj podgląd importu')
    if expiry < time.time() or not hmac.compare_digest(token, preview_token(secret, raw, fmt, actor['id'], expiry)):
        raise Problem(409, 'Podgląd wygasł lub plik został zmieniony')
    rows = parse_upload(raw, fmt)
    with store.connect(write=True) as con:
        prepared, problems = validate_rows(con, rows)
        if problems:
            raise Problem(409, 'Import zablokowany. Wykonaj ponownie podgląd: ' + problems[0]['detail'])
        ids = []
        for data, stock in prepared:
            item_id = store.add_item(data, actor, connection=con)
            # Zero is kept explicitly, with an audit entry; receive requires >0.
            store.stock_operation(dict(stock, item_id=item_id,
                action='receive' if stock['quantity'] else 'adjust', note='Import nowych pozycji'), actor, connection=con)
            ids.append(item_id)
    return dict(count=len(ids), item_ids=ids)
